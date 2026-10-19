"""
文档解析器 — 提取 PDF / 图片 / Word / TXT 中的文本内容
供 Agent 进行语义理解和推荐/规划

OCR（引擎探测、倾角校正、置信度尾注、两条降级腿）不在此文件，见 `app/agents/ocr_leg.py`
（2026-10-15 拆分：本文件曾被 deskew＋警示分级顶到 670 行，越过「文件 ≤600 行」硬门禁）。
"""
import os
import time
import zipfile
import asyncio
from ..core.logging import setup_logging
from ..core.stream_utils import UPLOAD_TEXT_MAX_CHARS
from . import ocr_leg

logger = setup_logging()

# ---------- 解析资源上限（2026-09-07 审查 P1）----------
# 20MB 大小限制约束不了页数/内嵌图数：1MB PDF 可含上万页/图，每页渲染 150dpi+
# OCR（全局锁串行）→ 单个恶意上传可占住 OCR 通道小时级并占满执行线程池。
# 超限截断并注明，不整体失败——正常文档远低于此阈值。
_MAX_PDF_PAGES = 200        # PDF 最大处理页数
_MAX_OCR_IMAGES = 50        # 单文档最大 OCR 图片数（内嵌图/整页渲染共用）
_PARSE_DEADLINE_S = 60.0    # 单文档解析 wall-clock 预算
# AG-1（2026-09-19 审查）：zip 炸弹防护——20MB 上传可声明 GB 级解压成员，
# 图片数与 deadline 都拦不住 z.read() 一次读爆内存；单成员解压后大小上限
# 80MB（手机原图 <20MB，留足冗余），超限成员跳过不落内存
_MAX_IMAGE_BYTES = 80 * 1024 * 1024
# ponytail: deadline 是协作式检查（页/图粒度），单次 OCR 无法抢占，
# 实际耗时最多超出预算一次单页 OCR；如需硬中断改子进程隔离。


def _read_member_limited(z: zipfile.ZipFile, name: str, limit: int) -> bytes | None:
    """AG-1（2026-09-19 审查）：流式限读 zip 成员，解压量超 limit 返回 None。

    file_size 是 zip 头部声明值，攻击者可谎报小值，只有按真实读取字节封顶
    才能防住真炸弹；返回超限标记让调用方决定跳过文案。
    """
    with z.open(name) as mf:
        data = mf.read(limit + 1)
    return None if len(data) > limit else data


# RAG-5（2026-09-20 审查）：AG-1 只封了 word/media 成员的真实读取，但
# DocxDocument() 构造即整体解压 XML 部件——高压缩比/超大声明解压尺寸仍可
# 一次性打爆 1GB 容器。解析前用 zipinfo 校验声明总尺寸与压缩比，超限拒绝。
_DOCX_MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
_DOCX_MAX_COMPRESSION_RATIO = 60  # 正常文本 docx 压缩比 <30，炸弹常见 >1000


def _docx_zip_bomb_reason(filepath: str) -> str:
    """命中封项返回拒绝原因（固定文案，不含路径），否则返回空串。"""
    try:
        with zipfile.ZipFile(filepath) as z:
            infos = z.infolist()
    except (zipfile.BadZipFile, OSError):
        return "文件不是有效的 docx 容器"
    total_uncompressed = sum(i.file_size for i in infos)
    if total_uncompressed > _DOCX_MAX_UNCOMPRESSED_BYTES:
        return "解压总大小超过上限"
    total_compressed = sum(i.compress_size for i in infos)
    if total_compressed and total_uncompressed > (
            total_compressed * _DOCX_MAX_COMPRESSION_RATIO):
        return "压缩比异常（疑似 zip bomb）"
    return ""


# ============================================================
# PDF 解析（PyMuPDF）— 文本 + 内嵌图片 OCR
# ============================================================
try:
    import fitz  # PyMuPDF
    _HAS_PDF = True
except ImportError:
    _HAS_PDF = False
    logger.warning("PyMuPDF 未安装，PDF 解析不可用")

def _truncation_note(pages: int, images: int, seconds: float, reason: str) -> str:
    """截断标注（拼在文本尾部，让上层与用户知道内容不完整）"""
    return (f"\n\n[文档过大已截断：处理了 {pages} 页 / {images} 张图 / {seconds:.0f} 秒，"
            f"原因：{reason}；其余内容未解析]")


def _clip_with_note(text: str, *, full_text: bool = False) -> str:
    """解析输出统一上限（与进 prompt 的消费口径**同一个数**）：超出截断并在头部标注。

    `full_text=True` 是离线清洗那条腿的口径：清洗件是拿去建索引的，不该按"一次塞进 prompt
    的额度"砍尾（10-18 实测：东莞那份 46431 字被 40000 字上限切在第 32 页，真值 137 条只剩
    104 条）。默认关＝现网一字未动。

    旧实现 `text[:150000]` 静默砍尾：超出 prompt 消费口径（4 万字）的部分零读取方，
    白耗 OCR 全局锁与 Redis；且截断对 LLM 与用户均不可见——LLM 拿前 4 万字答题，
    用户以为全文都在。标注放**头部**：下游按任意长度再切片都不会把它切掉。
    """
    if full_text or len(text) <= UPLOAD_TEXT_MAX_CHARS:
        return text
    return (f"[文档过长已截断：原文约 {len(text)} 字，以下仅保留前 "
            f"{UPLOAD_TEXT_MAX_CHARS} 字，其余内容未参与解析]\n"
            + text[:UPLOAD_TEXT_MAX_CHARS])


def _render_pages_ocr(filepath: str, deadline_s: float = _PARSE_DEADLINE_S) -> str:
    """整页渲染 + 逐页 OCR（纯扫描件兜底）。

    旧实现把 PDF 路径直接丢给图片 OCR 引擎（RapidOCR/Tesseract 都打不开 PDF），必然失败；
    这里用 fitz 把每页渲染成位图再走统一的 OCR 通道。
    受页数上限与 deadline 约束（2026-09-07 审查 P1）。`deadline_s` 的默认值绑的就是那个常量，
    现网调用因此与今天逐字节相同；只有离线清洗那条腿会显式声明更宽的预算（Bug日志 74 ①）。
    """
    parts = []
    rendered = 0  # 实际渲染页数（截断标注用：每页渲染即一张位图，原恒标 0 失真）
    started = time.monotonic()
    with fitz.open(filepath) as doc:
        for page_num, page in enumerate(doc):
            if page_num + 1 > _MAX_PDF_PAGES:
                parts.append(_truncation_note(_MAX_PDF_PAGES, rendered, time.monotonic() - started, "页数超限"))
                break
            if time.monotonic() - started > deadline_s:
                parts.append(_truncation_note(page_num, rendered, time.monotonic() - started, "解析超时"))
                break
            pix = page.get_pixmap(dpi=150)
            rendered += 1
            ocr_text = ocr_leg._ocr_image_bytes(pix.tobytes("png"), "png", f"第{page_num+1}页")
            if ocr_text and ocr_text.strip():
                parts.append(ocr_text)
    return "\n".join(parts)


async def parse_pdf(filepath: str, *, deadline_s: float = _PARSE_DEADLINE_S,
                full_text: bool = False) -> str:
    """解析 PDF：先提取文本，再提取内嵌图片进行 OCR，合并返回。

    `deadline_s` 是给这份件的 wall-clock 预算，默认就是 `_PARSE_DEADLINE_S`（现网上传解析因此
    一字不变）；离线清洗那条腿按批次显式声明，为的是别让 60 秒把整份件切成半份（Bug日志 74 ①）。
"""
    if not _HAS_PDF:
        return "[PDF 解析引擎未安装]"
    loop = asyncio.get_running_loop()

    def _extract_all():
        text_parts = []
        started = time.monotonic()
        page_count = 0
        img_count = 0
        stop_reason = ""
        with fitz.open(filepath) as doc:
            for page_num, page in enumerate(doc):
                # deadline 检查先于计数自增（2026-09-10 审查 P3：原顺序把未处理的
                # 当前页也计入 page_count，"处理了 N 页"标注虚高 1）
                if time.monotonic() - started > deadline_s:
                    stop_reason = "解析超时"
                    break
                page_count += 1
                if page_count > _MAX_PDF_PAGES:
                    stop_reason = "页数超限"
                    page_count = _MAX_PDF_PAGES
                    break
                # 1. 提取文本
                text = page.get_text()
                if text.strip():
                    text_parts.append(f"--- 第 {page_num+1} 页 ---\n{text}")

                # 2. 提取本页内嵌图片做 OCR（受图片总数上限约束）
                images_on_page = page.get_images(full=True)
                for img_idx, img_ref in enumerate(images_on_page):
                    img_count += 1
                    if img_count > _MAX_OCR_IMAGES:
                        stop_reason = "内嵌图片数超限"
                        break
                    xref = img_ref[0]  # 图片引用编号
                    try:
                        img_dict = doc.extract_image(xref)
                        # 字节直接交给 OCR 腿（腿内 Image.open + 同一置信度口径）；
                        # 单张图字节解码失败由腿里记日志并返回空串，这里只兜 extract_image
                        ocr_text = ocr_leg._ocr_image_bytes(
                            img_dict["image"], img_dict["ext"],
                            f"第{page_num+1}页-图{img_idx+1}")
                        if ocr_text and ocr_text.strip():
                            text_parts.append(ocr_text)
                    except Exception as img_err:
                        # 单张内嵌图片损坏不该废掉整份文档，跳过并留排查日志（P2 修复：不再裸 pass）
                        logger.debug(f"跳过无法解析的内嵌图片(第{page_num+1}页-图{img_idx+1}): {img_err}")
                if img_count > _MAX_OCR_IMAGES:
                    break

        if stop_reason:
            # 超限截断如实标注（内容不完整必须让调用方知道，2026-09-07 审查 P1）
            text_parts.append(_truncation_note(
                page_count, min(img_count, _MAX_OCR_IMAGES),
                time.monotonic() - started, stop_reason))
        return "\n\n".join(text_parts) if text_parts else ""

    try:
        text = await loop.run_in_executor(None, _extract_all)
        if not text.strip():
            # 纯扫描件（无文本层且无内嵌图片）→ fitz 整页渲染后 OCR
            logger.info("PDF 无直接文本且无内嵌图片，尝试整页渲染 OCR...")
            text = await loop.run_in_executor(None, _render_pages_ocr, filepath, deadline_s)
        return _clip_with_note(text, full_text=full_text)
    except Exception as e:
        logger.warning(f"PDF 解析失败: {e}")
        return f"[PDF 解析错误: {e}]"


# ============================================================
# 图片 OCR（RapidOCR 优先，Tesseract 降级）— 独立图片文件
# 引擎判据与两条降级腿都在 ocr_leg.ocr_image_file（拆分 2026-10-15）；
# Tesseract 可用性由 _tesseract_available() 首次使用时惰性探测（09-22 审阅 P2），
# 拆分后这次探测发生在线程池里而不是事件循环上，产出与文案一字未变。
async def parse_image(filepath: str) -> str:
    """对独立图片文件进行 OCR（RapidOCR 优先）"""
    loop = asyncio.get_running_loop()
    try:
        text = await loop.run_in_executor(None, ocr_leg.ocr_image_file, filepath)
        return _clip_with_note(text)
    except Exception as e:
        return f"[图片解析错误: {e}]"


# ============================================================
# Word 解析（python-docx）— 文本 + 内嵌图片 OCR
# ============================================================
try:
    from docx import Document as DocxDocument
    _HAS_DOCX = True
except ImportError:
    _HAS_DOCX = False
    logger.warning("python-docx 未安装，Word 解析不可用")

async def parse_docx(filepath: str) -> str:
    """解析 .docx：先提取文本/表格，再提取内嵌图片进行 OCR，合并返回"""
    if not _HAS_DOCX:
        return "[Word 解析引擎未安装]"
    loop = asyncio.get_running_loop()

    def _extract_all():
        # RAG-5：先验封 zip bomb，再交给 DocxDocument 整体解压
        reason = _docx_zip_bomb_reason(filepath)
        if reason:
            logger.warning(f"Word 文件被拒（RAG-5 防护）: {reason}")
            return f"[Word 解析失败: {reason}]"
        doc = DocxDocument(filepath)
        parts = []

        # 1. 提取段落文本
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        if paragraphs:
            parts.append("\n".join(paragraphs))

        # 2. 提取表格文本：折叠形态与清洗侧共用一把尺（`kbclean.type_layer.flatten_table`），
        #    每行都带列名。原先的 `" | ".join(cells)` 会把列名丢掉，进 prompt 就是
        #    「28 ｜ 张三」这种无上下文碎片，检索与问答都读不懂它出自哪张表哪一列。
        from app.services.kbclean.type_layer import flatten_table
        tables_text: list[str] = []
        for table in doc.tables:
            tables_text.extend(flatten_table([[cell.text for cell in row.cells] for row in table.rows]))
        if tables_text:
            parts.append("【表格内容】\n" + "\n".join(tables_text))

        # 3. 提取内嵌图片做 OCR（受图片数上限与 deadline 约束，2026-09-07 审查 P1）
        # .docx 本质是 ZIP，图片在 word/media/ 下
        ocr_results = []
        started = time.monotonic()
        try:
            with zipfile.ZipFile(filepath, 'r') as z:
                # Zip Slip 防护（P0 #4）：仅读取 word/media/ 下无路径穿越的成员。
                # 关键：必须用 RAW 成员名判断 — os.path.normpath 会把 'word/media/../X'
                # 折叠成 'word/X' 从而绕过 '..' 检查。ZIP 成员名恒用 '/' 分隔，按 '/' 切分即可。
                media_files = [
                    f for f in z.namelist()
                    if f.startswith('word/media/')
                    and not f.startswith('/')
                    and '..' not in f.split('/')
                ]
                for idx, media_path in enumerate(sorted(media_files)):
                    if idx + 1 > _MAX_OCR_IMAGES:
                        ocr_results.append(f"[截断：内嵌图片超过 {_MAX_OCR_IMAGES} 张，其余未解析]")
                        break
                    if time.monotonic() - started > _PARSE_DEADLINE_S:
                        ocr_results.append("[截断：图片 OCR 超时，其余未解析]")
                        break
                    # AG-1：真实解压字节封顶，超限成员不进内存
                    img_bytes = _read_member_limited(z, media_path, _MAX_IMAGE_BYTES)
                    if img_bytes is None:
                        ocr_results.append(
                            f"[跳过 {media_path}：解压后超过 {_MAX_IMAGE_BYTES // (1024 * 1024)}MB 上限]")
                        continue
                    ext = os.path.splitext(media_path)[1].lstrip('.')
                    label = f"文档插图{idx+1}"
                    ocr_text = ocr_leg._ocr_image_bytes(img_bytes, ext, label)
                    if ocr_text and ocr_text.strip():
                        ocr_results.append(ocr_text)
        except Exception as e:
            logger.warning(f"Word 图片提取失败: {e}")

        if ocr_results:
            parts.append("【文档内嵌图片文字】\n" + "\n".join(ocr_results))

        return "\n\n".join(parts) if parts else ""

    try:
        text = await loop.run_in_executor(None, _extract_all)
        return _clip_with_note(text)
    except Exception as e:
        logger.warning(f"Word 解析失败: {e}")
        return f"[Word 解析错误: {e}]"


# ============================================================
# TXT 解析（直接读取）
# ============================================================
async def parse_txt(filepath: str, encoding: str = None) -> str:
    """解析纯文本文件"""
    loop = asyncio.get_running_loop()
    def _read():
        encodings = [encoding, 'utf-8', 'gbk', 'gb2312', 'utf-16'] if not encoding else [encoding]
        for enc in encodings:
            if not enc:
                continue
            try:
                with open(filepath, 'r', encoding=enc) as f:
                    return f.read()
            except (UnicodeDecodeError, UnicodeError):
                continue
        return "[无法解码文件内容]"
    try:
        text = await loop.run_in_executor(None, _read)
        return _clip_with_note(text)
    except Exception as e:
        return f"[文本解析错误: {e}]"


# ============================================================
# 统一入口
# ============================================================
SUPPORTED_EXTENSIONS = {
    '.pdf':  'PDF 文档',
    '.jpg':  'JPEG 图片',
    '.jpeg': 'JPEG 图片',
    '.png':  'PNG 图片',
    '.bmp':  'BMP 图片',
    '.webp': 'WebP 图片',
    '.docx': 'Word 文档',
    '.doc':  'Word 文档（旧格式，需先转 .docx）',
    '.txt':  '纯文本',
    '.md':   'Markdown',
    '.csv':  'CSV 表格',
    '.json': 'JSON 数据',
    '.xml':  'XML 文档',
}

async def parse_document(filepath: str, filename: str) -> dict:
    """
    统一文档解析入口
    返回: {"text": str, "format": str, "pages": int, "success": bool}
    """
    ext = os.path.splitext(filename)[1].lower()
    logger.info(f"解析文档: {filename} (ext={ext})")

    if ext == '.pdf':
        text = await parse_pdf(filepath)
    elif ext in ('.jpg', '.jpeg', '.png', '.bmp', '.webp'):
        text = await parse_image(filepath)
    elif ext == '.docx':
        text = await parse_docx(filepath)
    elif ext in ('.txt', '.md', '.csv', '.json', '.xml'):
        text = await parse_txt(filepath)
    else:
        return {"text": "", "format": ext, "success": False, "error": f"不支持的文件格式: {ext}"}

    # 错误判定用显式标记而非 startswith("[")：正文以 "[" 开头的正常文档（如
    # JSON 数组导出、引用块）曾被误判为解析失败
    _ERROR_MARKERS = (
        "[PDF 解析错误", "[PDF 解析引擎未安装]",
        "[Word 解析错误", "[Word 解析引擎未安装]",
        "[图片解析错误", "[OCR 引擎未安装]", "[OCR 识别失败",
        "[文本解析错误", "[无法解码文件内容]",
    )
    is_error = any(text.startswith(m) for m in _ERROR_MARKERS)
    success = bool(text.strip()) and not is_error
    return {
        "text": text,
        "format": SUPPORTED_EXTENSIONS.get(ext, ext),
        "success": success,
        "error": text if is_error else None
    }
